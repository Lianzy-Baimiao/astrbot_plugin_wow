# -*- coding: utf-8 -*-
"""处罚名单查询：扫描目录下 xlsx 表格（角色名/服务器名两列），按角色/服务器匹配。

角色名脱敏规则（官方公告的写法）：**一个 `*` 恰好代表一个字符**，
所以 `张*` 只可能是 2 字名（张三），`张*丰` 才是 3 字名（张三丰），两者不互相命中。
查询时给完整名或同样带 `*` 的脱敏名都行，按「等长 + 逐位比对」匹配。
"""

from __future__ import annotations

import asyncio
import logging
import re
import threading
from pathlib import Path

logger = logging.getLogger("astrbot_plugin_wow.punish")

_lock = threading.Lock()
_cache: dict[str, tuple[int, int, list[dict]]] = {}

_ROLE_KEYS = ("角色名", "角色昵称", "角色", "昵称", "名字", "游戏名", "玩家")
_REALM_KEYS = ("服务器名", "服务器", "区服", "服务区")

# 全角星号：输出时替换半角 *，避免客户端把 `张**丰` 的 ** 渲染成 Markdown 加粗
_STAR_OUT = "＊"
_STARS = "*＊"  # 输入侧两种都认（用户可能直接复制上一条回复里的名字）

# 规范文件名：20260828_至暗之夜第二赛季_PVE_处罚名单.xlsx
# 第 3 段是「类别」，除 PVE/PVP 外还有 PVE淬炼英雄 / PVEPVP / 外挂欺诈 等写法，原样展示
_STEM = re.compile(r"^(\d{8})_([^_]+)_([^_]+)_处罚名单$")
# 旧文件名兜底：开头的 8 位或 4 位日期 + 可选 pve/pvp
_STEM_OLD = re.compile(r"^(\d{8}|\d{4})\s*(pve|pvp)?", re.I)


def _source(stem: str) -> dict:
    """文件名 -> {date, season, mode}；认不出的字段留空。"""
    m = _STEM.match(stem)
    if m:
        return {"date": m.group(1), "season": m.group(2), "mode": m.group(3)}
    m = _STEM_OLD.match(stem)
    if m:
        d = m.group(1)
        return {
            "date": d if len(d) == 8 else "2026" + d,
            "season": "",
            "mode": (m.group(2) or "").upper(),
        }
    return {"date": "", "season": "", "mode": ""}


def _pretty_date(date8: str) -> str:
    """20260828 -> 2026年8月28日；认不出原样返回。"""
    if len(date8) != 8 or not date8.isdigit():
        return date8
    return f"{date8[:4]}年{int(date8[4:6])}月{int(date8[6:]) }日"


def _match(real: str, name: str) -> bool:
    """名单里的脱敏名 real 与查询名 name 是否可能是同一个角色。

    一个 * = 一个字符，所以先要求**长度相等**，再逐位比对：
    任一侧该位是 * 就算通配，否则必须字符相同。
    这样 `张*`（2 字）不会命中 `张三丰`（3 字），`张*丰` 才会。
    """
    if len(real) != len(name):
        return False
    for a, b in zip(real, name):
        if a in _STARS or b in _STARS:
            continue
        if a != b:
            return False
    return True


def _detect_header(row: list) -> tuple[int, int]:
    """判断一行是否为表头，返回 (角色列索引, 服务器列索引)；非表头返回 (-1, -1)。"""
    role_idx = realm_idx = -1
    for i, cell in enumerate(row):
        s = str(cell or "").strip()
        if role_idx < 0 and s in _ROLE_KEYS:
            role_idx = i
        elif realm_idx < 0 and s in _REALM_KEYS:
            realm_idx = i
    return role_idx, realm_idx


def _load(xlsx: Path) -> list[dict]:
    """读取一个 xlsx：返回 [{name, realm}]。按 mtime+size 缓存，文件变了自动重读。

    通用表头识别：任一整行命中「角色/服务器」关键词即视为表头（支持多表合并的
    工作簿，表头位置不限第一行），其余行按当前列索引取值。
    """
    stat = xlsx.stat()
    key = str(xlsx)
    with _lock:
        hit = _cache.get(key)
        if hit and hit[0] == stat.st_mtime_ns and hit[1] == stat.st_size:
            return hit[2]

    from openpyxl import load_workbook

    rows: list[dict] = []
    role_idx = realm_idx = -1
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                if not r or not any(v is not None for v in r):
                    continue
                vals = [str(v).strip() if v is not None else "" for v in r]
                ri, re_ = _detect_header(vals)
                if ri >= 0:
                    role_idx, realm_idx = ri, re_
                    continue
                if role_idx < 0:
                    continue
                name = vals[role_idx] if role_idx < len(vals) else ""
                realm = vals[realm_idx] if 0 <= realm_idx < len(vals) else ""
                if name:
                    rows.append({"name": name, "realm": realm})
    finally:
        wb.close()

    with _lock:
        _cache[key] = (stat.st_mtime_ns, stat.st_size, rows)
    return rows


def invalidate_cache() -> None:
    """清空 xlsx 解析缓存（新名单落盘后调用，让下次查询读到新表）。"""
    with _lock:
        _cache.clear()


def query(punish_dir: Path, name: str, realm: str | None = None) -> list[dict]:
    """全目录扫描，返回命中记录（按公告日期倒序，最近的在前）。

    每条：{date, season, mode, name, realm, file}
    name 可为完整名（张三丰）或脱敏名（张**丰）；realm 精确匹配。
    """
    if not punish_dir.is_dir():
        raise RuntimeError(f"处罚名单目录不存在：{punish_dir}")
    out: list[dict] = []
    # rglob：名单按赛季归档到子目录（如「归档/」）后仍要能查到
    xlsx_list = sorted(p for p in punish_dir.rglob("*.xlsx") if not p.name.startswith("~$"))
    logger.info("[punish] 扫描 %s，共 %d 个 xlsx", punish_dir, len(xlsx_list))
    for xlsx in xlsx_list:
        try:
            rows = _load(xlsx)
            src = _source(xlsx.stem)
            for row in rows:
                if realm is not None and row["realm"] != realm:
                    continue
                if _match(name, row["name"]):
                    out.append({**src, "name": row["name"], "realm": row["realm"],
                                "file": xlsx.name})
        except Exception as e:  # noqa: BLE001
            logger.warning("处罚名单读取失败 %s: %s", xlsx.name, e)
    out.sort(key=lambda h: (h["date"], h["mode"]), reverse=True)
    logger.info("[punish] 查询「%s」@%s，命中 %d 条", name, realm or "*", len(out))
    return out


async def query_async(punish_dir: Path, name: str, realm: str | None = None) -> list[dict]:
    """query 的线程版：首轮要读 38 万行 xlsx（约 10s），不能阻塞事件循环。"""
    return await asyncio.to_thread(query, punish_dir, name, realm)


MAX_LINES = 30  # 高频脱敏名（如「丨****丨」有 200+ 条）会把群消息刷爆，超出只列最近的


def safe_name(name: str) -> str:
    """半角 * 换成全角＊：客户端按 Markdown 渲染时，`张**丰` 的 ** 会被吃成加粗。"""
    return name.replace("*", _STAR_OUT)


def format_hits(name: str, realm: str | None, hits: list[dict]) -> str:
    """命中记录 -> 群消息文案（按赛季分组，日期从近到远）。"""
    where = f"·{realm}" if realm else ""
    lines = [f"「{safe_name(name)}{where}」共查到 {len(hits)} 次处罚记录", ""]
    shown = hits[:MAX_LINES]
    last_season = object()
    for h in shown:
        season = h["season"] or "未知赛季"
        if season != last_season:
            lines.append(f"【{season}】")
            last_season = season
        mode = f"{h['mode']} " if h["mode"] else ""
        date = _pretty_date(h["date"]) if h["date"] else "日期不详"
        lines.append(f"  · {date} {mode}名单：{safe_name(h['name'])}"
                     f"（{h['realm'] or '服务器未记录'}）")
    if len(hits) > len(shown):
        lines.append(f"  …… 另有 {len(hits) - len(shown)} 条较早记录未列出")
    lines.append("")
    lines.append("名单为官方公示的脱敏名，同名玩家可能重复，仅供参考。")
    return "\n".join(lines)
