# -*- coding: utf-8 -*-
"""杂项服务：吃什么/天赋/物价/低保/语录/地下堡。"""

from __future__ import annotations

import asyncio
import datetime as dt
import http.cookiejar
import logging
import random
import re
import urllib.error
import urllib.parse
import urllib.request

from ..const import UA
from ..data.delvers import day_number, get_menu
from ..data.foods import pick_menu
from ..data.quotes import pick_quote
from ..data.talents import translate_talent_key
from ..store import load_json, save_json

logger = logging.getLogger("astrbot_plugin_wow.misc")

# ---------------------------- 吃什么 ----------------------------


def what_to_eat() -> str:
    """按原版 chishenme getMenu()：4-10 点抽早餐 0~2 道，其余时段午/晚混抽 0~4 道。"""
    return pick_menu()


# ---------------------------- 天赋 ----------------------------

ARCHON_SLUG = {
    "血DK": "blood/death-knight", "冰DK": "frost/death-knight", "邪DK": "unholy/death-knight",
    "噬灭": "devourer/demon-hunter", "浩劫": "havoc/demon-hunter", "复仇": "vengeance/demon-hunter",
    "平衡": "balance/druid", "野德": "feral/druid", "熊德": "guardian/druid", "奶德": "restoration/druid",
    "增辉": "augmentation/evoker", "湮灭": "devastation/evoker", "恩护": "preservation/evoker",
    "兽王": "beast-mastery/hunter", "射击": "marksmanship/hunter", "生存": "survival/hunter",
    "奥法": "arcane/mage", "火法": "fire/mage", "冰法": "frost/mage",
    "酒仙": "brewmaster/monk", "织雾": "mistweaver/monk", "踏风": "windwalker/monk",
    "奶骑": "holy/paladin", "防骑": "protection/paladin", "惩戒": "retribution/paladin",
    "戒律": "discipline/priest", "神牧": "holy/priest", "暗牧": "shadow/priest",
    "刺杀": "assassination/rogue", "狂徒": "outlaw/rogue", "敏锐": "subtlety/rogue",
    "元素": "elemental/shaman", "增强": "enhancement/shaman", "恢复": "restoration/shaman",
    "痛苦": "affliction/warlock", "恶魔": "demonology/warlock", "毁灭": "destruction/warlock",
    "武器": "arms/warrior", "狂暴": "fury/warrior", "防战": "protection/warrior",
}

SPEC_LIST = {
    "死亡骑士": ["血DK", "冰DK", "邪DK"], "恶魔猎手": ["噬灭", "浩劫", "复仇"],
    "德鲁伊": ["平衡", "野德", "熊德", "奶德"], "唤魔师": ["增辉", "湮灭", "恩护"],
    "猎人": ["兽王", "射击", "生存"], "法师": ["奥法", "火法", "冰法"],
    "武僧": ["酒仙", "织雾", "踏风"], "圣骑士": ["奶骑", "防骑", "惩戒"],
    "牧师": ["戒律", "神牧", "暗牧"], "潜行者": ["刺杀", "狂徒", "敏锐"],
    "萨满": ["元素", "增强", "恢复"], "术士": ["痛苦", "恶魔", "毁灭"],
    "战士": ["武器", "狂暴", "防战"],
}

_TALENT_RE = re.compile(r"calc/blizzard/([A-Za-z0-9+/=]+)")
_TALENT_RE2 = re.compile(r"([C][A-Za-z0-9+/=]{80,})")
# 页面内嵌十几个 data:image/...;base64 图片，字符集跟天赋串一样，会被兜底正则
# 误当成天赋码返回，搜兜底之前先剔掉。
_DATA_URI_RE = re.compile(r"data:[^;'\"\s]+;base64,[A-Za-z0-9+/=]+")

# Archon.gg 有两层拦截：Cloudflare 机器人检查，过了之后还有站点自建的
# Human Verification 表单页（HTTP 200，所以 raise_for_status 发现不了）。
# httpx 过不了第一层（403 "Just a moment..."），标准库 urllib 可以，所以这里
# 不走 net.py 的共享客户端。第二层把页面里的隐藏字段原样 POST 回去就能过，
# 响应体直接就是目标页面，拿到的 human_verified cookie 服务端给约 90 天。
_CHALLENGE_URL = "https://www.archon.gg/human-challenge"
_CHALLENGE_MARK = 'action="/human-challenge"'
_INPUT_RE = re.compile(r"<input\b[^>]*>", re.I)
_ATTR_RE = re.compile(r'(\w+)\s*=\s*"([^"]*)"')

_archon_op: urllib.request.OpenerDirector | None = None


def _archon_opener() -> urllib.request.OpenerDirector:
    """带 CookieJar 的 opener（惰性创建）：human_verified cookie 在进程内复用，
    重启后重新解一次门（多一个 POST，约 0.4 秒）。"""
    global _archon_op
    if _archon_op is None:
        jar = http.cookiejar.CookieJar()
        _archon_op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        _archon_op.addheaders = [("User-Agent", UA)]
    return _archon_op


def _hidden_fields(page: str) -> dict[str, str]:
    """取表单里的隐藏字段（不依赖属性顺序）。"""
    out: dict[str, str] = {}
    for tag in _INPUT_RE.findall(page):
        attrs = dict(_ATTR_RE.findall(tag))
        if attrs.get("type") == "hidden" and "name" in attrs:
            out[attrs["name"]] = attrs.get("value", "")
    return out


def _archon_get_sync(url: str, timeout: float) -> str:
    """GET 页面；撞上人机校验就提交表单，POST 的响应体即目标页面。阻塞，交给线程池跑。"""
    op = _archon_opener()
    page = op.open(url, timeout=timeout).read().decode("utf-8", "replace")
    if _CHALLENGE_MARK not in page:
        return page
    fields = _hidden_fields(page)
    if "signature" not in fields:
        logger.warning("Archon.gg 人机校验表单字段变了，无法自动通过：%s", sorted(fields))
        return ""
    req = urllib.request.Request(
        _CHALLENGE_URL,
        data=urllib.parse.urlencode(fields).encode(),
        headers={"Content-Type": "application/x-www-form-urlencoded", "Referer": url},
    )
    return op.open(req, timeout=timeout).read().decode("utf-8", "replace")


async def _fetch_archon_talent(url: str) -> str:
    try:
        page = await asyncio.to_thread(_archon_get_sync, url, 15.0)
    except urllib.error.HTTPError as e:
        if e.code == 403:
            logger.warning("Archon.gg 被 Cloudflare 拦截（403，未过机器人检查）：%s", url)
        else:
            logger.warning("Archon.gg 抓取失败 %s: HTTP %s", url, e.code)
        return ""
    except Exception as e:  # noqa: BLE001
        logger.warning("Archon.gg 抓取失败 %s: %s", url, e)
        return ""
    m = _TALENT_RE.search(page)
    if m:
        return m.group(1)
    m2 = _TALENT_RE2.search(_DATA_URI_RE.sub("", page))
    if m2:
        logger.info("Archon.gg 页面无 calc/blizzard/，改用兜底正则：%s", url)
        return m2.group(1)
    if _CHALLENGE_MARK in page:
        logger.warning("Archon.gg 人机校验未通过，仍停在校验页：%s", url)
    return ""


def is_talent_key(key: str) -> bool:
    """能否识别为职业或专精（免前缀触发时用来过滤正常聊天）。"""
    k = translate_talent_key(key.strip())
    return k in SPEC_LIST or k in ARCHON_SLUG


async def talent_info(key: str) -> str:
    key = translate_talent_key(key.strip())
    if key in SPEC_LIST:
        return f"{key}天赋推荐，请输入具体天赋名称获取：\n" + "\n".join(f"  {s}" for s in SPEC_LIST[key])
    slug = ARCHON_SLUG.get(key)
    if not slug:
        return "未找到匹配的职业或天赋：" + key
    raid_url = f"https://www.archon.gg/wow/builds/{slug}/raid/overview/mythic/all-bosses/this-week"
    mp_url = f"https://www.archon.gg/wow/builds/{slug}/mythic-plus/overview/10/all-dungeons/this-week"
    raid, mp = await _fetch_archon_talent(raid_url), await _fetch_archon_talent(mp_url)
    if not raid and not mp:
        return f"{key} 天赋获取失败：Archon.gg 没返回数据（暂无数据或被拦截，详见日志）"
    lines = [f"{key} 天赋推荐（Archon.gg）"]
    if raid:
        lines.append("团：" + raid)
    if mp:
        lines.append("秘：" + mp)
    return "\n".join(lines)


# ---------------------------- 物价（本地 xlsx） ----------------------------

def _cell(row: tuple, idx: int):
    return row[idx] if len(row) > idx else None


def _query_price_sync(items: list[str]) -> str:
    """xlsx 读取是同步阻塞的，交给线程池跑（B=价格 C=名称 D=装等 E=数量）。"""
    from openpyxl import load_workbook
    from ..store import data_dir
    xlsx = data_dir() / "prices.xlsx"
    if not xlsx.exists():
        return "物价表不存在，请将 CustomDecode xlsx 放到插件数据目录并命名为 prices.xlsx"
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        ws = wb.active
        out = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = _cell(row, 2)
            if name is None:
                continue
            for q in items:
                q = q.strip()
                if q and q in str(name):
                    out.append((str(name), _cell(row, 1), _cell(row, 3), _cell(row, 4)))
                    break
            if len(out) >= 5:
                break
    finally:
        wb.close()
    if not out:
        return "未找到相关物品，请检查物品名"
    lines = [f"查询 {len(items)} 种物品，结果如下："]
    for name, price, ilvl, qty in out:
        try:
            price_s = f"{float(price):,}"
        except (TypeError, ValueError):
            price_s = str(price) if price is not None else "?"
        lines.append(f"{name}：{price_s} 金（装等 {ilvl or '?'}，数量 {qty or '?'}）")
    return "\n".join(lines)


async def query_price(items: list[str]) -> str:
    """查询本地物价表（CustomDecode xlsx）。"""
    return await asyncio.to_thread(_query_price_sync, items)


# ---------------------------- 低保（每日运势） ----------------------------

def daily_fortune(uid: str) -> str:
    today = dt.date.today().isoformat()
    data = load_json("fortune.json", {}) or {}
    if data.get(str(uid)) == today:
        return "今天已经签到过低保啦，明天再来！"
    luck = random.choice(["欧皇附体", "小赚一笔", "平平淡淡", "非酋降临"])
    data[str(uid)] = today
    save_json("fortune.json", data)
    msg = {
        "欧皇附体": "今日低保运势：欧皇附体！开箱子必出传说！",
        "小赚一笔": "今日低保运势：小赚一笔，饰品提升+1。",
        "平平淡淡": "今日低保运势：平平淡淡，就当无事发生。",
        "非酋降临": "今日低保运势：非酋降临，建议今天别开箱。",
    }
    return msg[luck]


# ---------------------------- 语录 / 地下堡 ----------------------------

def quote_card(arg: str = "") -> dict:
    return pick_quote(arg)


def delver_menu() -> str:
    now = dt.datetime.now(dt.timezone.utc)
    day = day_number(now)
    items = get_menu(now)
    local = dt.datetime.now()
    date_str = local.strftime("%Y年%m月%d日")
    lines = [f"📅 {date_str}（周期第{day}天）可获取饰品："]
    lines.extend(f"{i + 1}. {name}" for i, name in enumerate(items))
    lines.append("提示：每5天为一个周期循环刷新")
    return "\n".join(lines)